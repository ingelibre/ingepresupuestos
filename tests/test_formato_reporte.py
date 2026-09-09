# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Marco Sumari
# This file is part of IngePresupuestos — https://ingepresupuestos.com
# Software libre bajo la GNU GPL v3 o posterior. Ver el archivo LICENSE.
"""Encabezado y pie configurables del PDF.

Corre con:  QT_QPA_PLATFORM=offscreen venv/bin/python3 tests/test_formato_reporte.py

Pedido de David Ramos (5 sep 2026): «activar, desactivar o personalizar algún
contenido de los encabezados superiores, y los pies de página (por ejemplo en
caso en la parte inferior solo se muestre los números de página)».

Dejar un texto en blanco significaba «usa el valor por defecto», así que no
había forma de pedir que un hueco quedara vacío — David escribió un punto en
el pie central para conseguirlo. Estas banderas son ese «nada» explícito.

Lo que se fija acá:

* `_oculto` solo acepta '1' como sí (las claves de formato son cadenas);
* cada ranura del pie desaparece por su cuenta, y las tres juntas dejan el
  pie en blanco;
* apagar el encabezado lo quita Y sube el cuerpo, que es de dónde sale el
  ahorro de hojas;
* los márgenes (`rep_margen_*`, en mm; 7 sep 2026) mueven de verdad el
  cuerpo, el encabezado y el pie, y con los valores por defecto la geometría
  es la de siempre.

Usa una copia temporal del seed, nunca la BD activa.
"""
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

from PySide6.QtWidgets import QApplication

_app = QApplication.instance() or QApplication([])

import core.config as cfg
import core.database as d

SEED = os.path.join(os.path.dirname(__file__), '..', 'presupuestos_seed.db')
_fd, _tmpdb = tempfile.mkstemp(suffix='_formato_test.db')
os.close(_fd)
shutil.copy(SEED, _tmpdb)
d.DB_PATH = _tmpdb
cfg.DB_PATH = _tmpdb
d.init_db()

import core.pdf_reports as pr   # noqa: E402  (después de fijar la BD)

_conn = d.get_db()
PID = _conn.execute(
    "SELECT proyecto_id FROM partidas GROUP BY proyecto_id "
    "HAVING COUNT(*) BETWEEN 40 AND 200 LIMIT 1"
).fetchone()[0]
# Un cliente para que el pie izquierdo tenga algo que imprimir.
_conn.execute("UPDATE proyectos SET cliente='MUNICIPALIDAD DE PRUEBA' WHERE id=?", (PID,))
_conn.commit()
_conn.close()

_tmpfiles = []

BANDERAS = ('rep_pie_izq_oculto', 'rep_pie_cen_oculto', 'rep_pie_der_oculto',
            'rep_encabezado_oculto', 'rep_pie_oculto')


# ── Andamio ──────────────────────────────────────────────────────────────────

def _flags(**kw):
    """Pone TODAS las banderas ('0' por defecto) y deja las pedidas en '1'."""
    for k in BANDERAS:
        d.set_config(k, '1' if kw.get(k) else '0')


def _pagina(n=1):
    """Palabras de la página `n` del Presupuesto, con su posición vertical.

    Se usa la 2ª página: la 1ª es la portada y no lleva encabezado ni pie.
    """
    import pdfplumber
    fd, path = tempfile.mkstemp(suffix='_fmt.pdf')
    os.close(fd)
    _tmpfiles.append(path)
    pr.generar_pdf_archivo('presupuesto', PID, path)
    with pdfplumber.open(path) as pdf:
        pg = pdf.pages[min(n, len(pdf.pages) - 1)]
        pie = ' '.join(w['text'] for w in
                       pg.crop((0, pg.height - 30, pg.width, pg.height)).extract_words())
        # Dónde empieza el CUERPO: la cabecera de la tabla del presupuesto.
        # No sirve la palabra más alta de la página — con encabezado, esa
        # palabra ES el encabezado.
        tops = [w['top'] for w in pg.extract_words() if w['text'] == 'Descripción']
        top = min(tops) if tops else None
    return pie, top


# ── Tests ────────────────────────────────────────────────────────────────────

def test_oculto_solo_acepta_uno():
    """Las claves de formato son cadenas; cualquier cosa que no sea '1' es no."""
    for valor in ('1',):
        assert pr._oculto({'k': valor}, 'k'), valor
    for valor in ('0', '', None, 'si', 'true', '2'):
        assert not pr._oculto({'k': valor}, 'k'), valor
    assert not pr._oculto({}, 'k')          # clave ausente
    assert not pr._oculto(None, 'k')        # formato ausente


def test_cada_ranura_del_pie_se_apaga_sola():
    _flags()
    completo, _ = _pagina()
    assert 'Cliente' in completo, completo
    assert 'Página' in completo, completo

    _flags(rep_pie_izq_oculto=True)
    sin_izq, _ = _pagina()
    assert 'Cliente' not in sin_izq, sin_izq
    assert 'Página' in sin_izq, sin_izq

    _flags(rep_pie_der_oculto=True)
    sin_der, _ = _pagina()
    assert 'Página' not in sin_der, sin_der
    assert 'Cliente' in sin_der, sin_der


def test_solo_el_numero_de_pagina():
    """El ejemplo textual de David: apagar izquierda y centro deja el número."""
    _flags(rep_pie_izq_oculto=True, rep_pie_cen_oculto=True)
    pie, _ = _pagina()
    assert pie.startswith('Página'), pie
    assert 'Cliente' not in pie, pie


def test_las_tres_banderas_dejan_el_pie_en_blanco():
    _flags(rep_pie_izq_oculto=True, rep_pie_cen_oculto=True, rep_pie_der_oculto=True)
    pie, _ = _pagina()
    assert pie == '', pie


def test_pie_apagado_no_imprime_nada_abajo():
    """Un solo interruptor para todo el pie (Marco, 8 sep 2026), aparte de
    las tres ranuras: apagado, abajo no queda ni el número de página."""
    _flags(rep_pie_oculto=True)
    pie, _ = _pagina()
    assert pie == '', pie
    _flags()
    pie, _ = _pagina()
    assert 'Página' in pie, pie


def test_sin_encabezado_el_cuerpo_sube():
    """Apagarlo no basta con no dibujarlo: la franja tiene que quedar para el
    cuerpo, que es de dónde sale el ahorro de hojas."""
    _flags()
    _, top_con = _pagina()

    _flags(rep_encabezado_oculto=True)
    _, top_sin = _pagina()

    assert top_con is not None and top_sin is not None, (top_con, top_sin)
    assert top_sin < top_con - 20, (top_sin, top_con)
    # Y el pie sigue intacto: son opciones independientes.
    pie_sin, _ = _pagina()
    assert 'Página' in pie_sin, pie_sin


# ── Márgenes ─────────────────────────────────────────────────────────────────
# Pedido de David Ramos (7 sep 2026): «una opción para configurar los
# márgenes». Cuatro claves en mm; con los valores por defecto la geometría es
# la de siempre.

def _margenes(**mm):
    for lado in ('sup', 'inf', 'izq', 'der'):
        d.set_config(f'rep_margen_{lado}',
                     str(mm.get(lado, pr.FORMATO_CLAVES[f'rep_margen_{lado}'])))


def _cuerpo_presupuesto():
    """(x del primer «Ítem», top de «Descripción», y del pie) en la 2ª página,
    en puntos PDF."""
    import pdfplumber
    fd, path = tempfile.mkstemp(suffix='_margen.pdf')
    os.close(fd)
    _tmpfiles.append(path)
    pr.generar_pdf_archivo('presupuesto', PID, path)
    with pdfplumber.open(path) as pdf:
        pg = pdf.pages[1]
        palabras = pg.extract_words()
        x_item = min(w['x0'] for w in palabras if w['text'] == 'Ítem')
        top = min(w['top'] for w in palabras if w['text'] == 'Descripción')
        pie = [w for w in palabras if w['text'] == 'Página']
        y_pie = pie[0]['top'] if pie else None
        ancho = pg.width
    return x_item, top, y_pie, ancho


def test_margenes_mm_acota_y_cae_al_defecto():
    m = pr.margenes_mm({})
    assert m == {'sup': 15.0, 'inf': 18.0, 'izq': 15.0, 'der': 15.0}, m
    m = pr.margenes_mm({'rep_margen_sup': 'abc', 'rep_margen_inf': '99',
                        'rep_margen_izq': '1', 'rep_margen_der': '20,5'})
    assert m['sup'] == 15.0, m          # ilegible → defecto
    assert m['inf'] == pr.MARGEN_MAX_MM  # por encima del tope
    assert m['izq'] == pr.MARGEN_MIN_MM  # por debajo del piso
    assert m['der'] == 20.5              # coma decimal
    for k in pr.MARGENES_CLAVES:
        assert k in pr.FORMATO_CLAVES, k


def test_con_los_margenes_por_defecto_la_geometria_es_la_de_siempre():
    """57 px laterales y el cuerpo a 100 px con encabezado — lo que el PDF
    tuvo siempre en A4 a 96 dpi. Si esto cambia, todos los PDF cambian."""
    _flags()          # el test anterior deja el encabezado apagado
    _margenes()
    r = pr._PdfRenderer({'nombre': 'x'}, 'x')
    assert (r.margin_x, r.margin_r) == (57, 57), (r.margin_x, r.margin_r)
    assert r.margin_top_body == 100, r.margin_top_body
    assert (r.header_dy, r.footer_dy) == (0, 0), (r.header_dy, r.footer_dy)


def test_el_margen_izquierdo_corre_el_cuerpo():
    _margenes()
    x0, _, _, _ = _cuerpo_presupuesto()
    _margenes(izq=30)
    x1, _, _, _ = _cuerpo_presupuesto()
    # 15 mm más = 42.5 pt más a la derecha (±2 pt de redondeo a píxel)
    assert abs((x1 - x0) - 15 / 25.4 * 72) < 2.5, (x0, x1)


def test_el_margen_superior_baja_el_cuerpo_y_el_inferior_sube_el_pie():
    _margenes()
    _, top0, pie0, _ = _cuerpo_presupuesto()
    _margenes(sup=30, inf=30)
    _, top1, pie1, _ = _cuerpo_presupuesto()
    assert abs((top1 - top0) - 15 / 25.4 * 72) < 2.5, (top0, top1)
    assert pie0 is not None and pie1 is not None, (pie0, pie1)
    assert abs((pie0 - pie1) - 12 / 25.4 * 72) < 2.5, (pie0, pie1)


def test_el_margen_derecho_no_deja_tinta_en_su_franja():
    """Con 35 mm a la derecha, la franja derecha del cuerpo queda en blanco:
    las tablas van al 100 % del ancho útil y tienen que encogerse con él."""
    from PySide6.QtCore import QSize
    from PySide6.QtPdf import QPdfDocument
    _margenes(der=35)
    fd, path = tempfile.mkstemp(suffix='_margen_der.pdf')
    os.close(fd)
    _tmpfiles.append(path)
    pr.generar_pdf_archivo('presupuesto', PID, path, with_cover=False)
    doc = QPdfDocument()
    doc.load(path)
    r = pr._PdfRenderer({'nombre': 'x'}, 'x')
    k = 1.5
    pt = 72.0 / r.dpi
    mr, top, bot = r.margin_r * pt * k, r.margin_top_body * pt * k, r.margin_bot_body * pt * k
    for i in range(min(3, doc.pageCount())):
        sz = doc.pagePointSize(i)
        img = doc.render(i, QSize(int(sz.width() * k), int(sz.height() * k)))
        w, h = img.width(), img.height()
        for y in range(int(top + 2), int(h - bot - 2), 2):
            for x in range(int(w - mr + 2), w, 2):
                c = img.pixelColor(x, y)
                assert not (c.alpha() > 0 and min(c.red(), c.green(), c.blue()) < 250), \
                    f"tinta en el margen derecho, página {i + 1} ({x}, {y})"
    _margenes()


# ── Esquemas de colores de títulos ───────────────────────────────────────────
# Pedido de David Ramos (5 sep 2026): color de títulos configurable y
# guardable. Diseño de Marco (8 sep): esquemas de fábrica + propios. Solo
# afecta a los reportes; «Clásico» reproduce los colores de siempre.

def _esquema(activo=None, propios=None):
    import json
    d.set_config('rep_esquema_titulos', activo or '')
    d.set_config('rep_esquemas_titulos', json.dumps(propios) if propios else '')


def test_clasico_es_el_color_de_siempre():
    """Clásico = theme.NIVEL_FG: la pantalla y el papel coinciden por defecto,
    y una instalación sin las claves imprime igual que antes."""
    from utils.theme import NIVEL_FG, NIVEL_MAX
    _esquema()
    cols = pr.colores_titulos()
    N = NIVEL_MAX
    assert N == pr.N_NIVELES == 9
    assert {n: cols[n] for n in range(1, N + 1)} == {n: NIVEL_FG[n].upper() for n in range(1, N + 1)}, cols
    assert cols[0] == '#1F2A38', cols      # sub-presupuesto: el slate-800 de siempre del PDF
    assert pr.ESQUEMAS_FABRICA['clasico']['colores'] == [NIVEL_FG[n] for n in range(1, N + 1)]
    assert all(len(v['colores']) == N for v in pr.ESQUEMAS_FABRICA.values())


def test_esquema_activo_desconocido_cae_a_clasico():
    _esquema(activo='no_existe')
    assert pr.colores_titulos() == pr.colores_titulos({'rep_esquema_titulos': 'clasico'})
    _esquema(activo='monocromo')
    assert set(pr.colores_titulos().values()) == {'#000000'}


def test_esquemas_propios_se_leen_y_lo_roto_se_ignora():
    propios = {
        'mio': {'nombre': 'Mío', 'colores': ['#111111', '#222222', 'rojo', None, '#555555']},
        'basura': 'no es un dict',
        'clasico': {'nombre': 'pisa fábrica', 'colores': ['#000000'] * 5},
    }
    _esquema(activo='mio', propios=propios)
    todos = pr.esquemas_titulos()
    assert 'basura' not in todos
    assert todos['clasico']['fabrica'] and todos['clasico']['colores'][0] == '#B71C1C'  # no se pisa
    cols = pr.colores_titulos()
    assert cols[1] == '#111111' and cols[2] == '#222222' and cols[5] == '#555555', cols
    assert cols[3] == '#6A1B9A' and cols[4] == '#AD1457', cols   # inválidos → Clásico
    assert cols[0] == '#1F2A38', cols                              # sin 'sub' → el de Clásico
    # Guardado con cinco colores (antes del 9 sep 2026): del 6 al 9 toma
    # los de Clásico, así el reporte de un proyecto hondo no se queda sin color.
    from utils.theme import NIVEL_FG
    assert {n: cols[n] for n in range(6, 10)} == {n: NIVEL_FG[n].upper() for n in range(6, 10)}, cols
    assert len(todos['mio']['colores']) == pr.N_NIVELES
    # JSON ilegible: como si no hubiera propios
    d.set_config('rep_esquemas_titulos', '{no json')
    assert all(v['fabrica'] for v in pr.esquemas_titulos().values())
    _esquema()


def test_el_esquema_cambia_el_color_del_titulo_en_el_pdf():
    """Con Monocromo, el título de nivel 1 del Presupuesto sale negro; con
    Clásico, rojo. Se lee el color de relleno de las letras con pdfplumber."""
    import pdfplumber

    def _color_titulo():
        fd, path = tempfile.mkstemp(suffix='_esq.pdf')
        os.close(fd)
        _tmpfiles.append(path)
        pr.generar_pdf_archivo('presupuesto', PID, path, with_cover=False)
        with pdfplumber.open(path) as pdf:
            pg = pdf.pages[0]
            # La primera palabra en mayúsculas y negrita bajo la cabecera de
            # la tabla es el título de nivel 1.
            for ch in pg.chars:
                if ch['text'].isalpha() and ch['text'].isupper() and 'Bold' in (ch.get('fontname') or '') \
                        and ch['top'] > 120:
                    c = ch.get('non_stroking_color')
                    return tuple(round(float(x), 2) for x in (c if isinstance(c, (list, tuple)) else [c]))
        return None

    _esquema(activo='clasico')
    rojo = _color_titulo()
    _esquema(activo='monocromo')
    negro = _color_titulo()
    _esquema()
    assert rojo is not None and negro is not None, (rojo, negro)
    assert negro in ((0, 0, 0), (0.0,)), negro
    assert rojo != negro, (rojo, negro)


def test_las_banderas_nuevas_estan_en_formato_claves():
    """Si no están, `get_formato` no las devuelve y el diálogo las pierde."""
    for k in BANDERAS:
        assert k in pr.FORMATO_CLAVES, k
        assert pr.FORMATO_CLAVES[k] == '0', k   # por defecto, todo se imprime


if __name__ == "__main__":
    fallos = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            try:
                fn()
                print(f"  OK  {name}")
            except AssertionError as e:
                fallos += 1
                print(f"  FAIL {name}: {e}")
    for f in _tmpfiles:
        if os.path.exists(f):
            os.unlink(f)
    if os.path.exists(_tmpdb):
        os.unlink(_tmpdb)
    sys.exit(1 if fallos else 0)


def test_los_nueve_colores_de_nivel_tienen_un_solo_dueno():
    """Árbol del presupuesto, metrados, control de obra, Gantt y el esquema
    Clásico de los reportes pintan los títulos con `theme.NIVEL_FG`. Hasta
    el 9 sep 2026 metrados y control de obra llevaban su copia (de 5 y de 4)
    y del 6 en adelante todo repetía el ámbar (reporte de David Ramos)."""
    from utils.theme import NIVEL_FG, NIVEL_MAX, nivel_fg
    import views.proyecto_view as PV
    import views.metrados_view as MV
    import views.control_obra_view as CO
    assert NIVEL_MAX == 9 and len(set(NIVEL_FG.values())) == 9   # nueve, distintos
    assert nivel_fg(10) == nivel_fg(9) == NIVEL_FG[9]           # se acota, no cae a negro
    assert {n: PV.NIVEL_ESTILO[n][0] for n in NIVEL_FG} == NIVEL_FG
    assert MV.NIVEL_COL is NIVEL_FG
    assert {n: CO.NIVEL_ESTILO[n][0] for n in NIVEL_FG} == NIVEL_FG
    # El CSS del PDF y el Excel tienen una clase/fuente para cada nivel.
    css = pr._base_css()
    for n in range(1, NIVEL_MAX + 1):
        assert f"tr.titulo{n} td" in css, n


def test_el_gantt_obedece_al_formato_de_encabezado_pie_y_leyenda():
    """Marco, 9 sep 2026: apagó encabezado y pie en «Editar formato» y el
    PDF del Gantt los seguía imprimiendo (el Centro y Ctrl+P le pasaban
    siempre True). Ahora `gantt_flags_desde_formato` los deriva de las
    mismas casillas; la leyenda tiene la suya (`rep_gantt_leyenda_oculta`)."""
    f = pr.gantt_flags_desde_formato({})
    assert f == {'incluir_header': True, 'incluir_footer': True,
                 'incluir_page': True, 'incluir_legend': True}
    f = pr.gantt_flags_desde_formato({'rep_encabezado_oculto': '1', 'rep_pie_oculto': '1'})
    assert not f['incluir_header'] and not f['incluir_footer'] and not f['incluir_page']
    assert f['incluir_legend']                      # la leyenda va aparte
    # Pie encendido pero ranura derecha vacía → sin número de página
    f = pr.gantt_flags_desde_formato({'rep_pie_der_oculto': '1'})
    assert f['incluir_footer'] and not f['incluir_page']
    assert not pr.gantt_flags_desde_formato({'rep_gantt_leyenda_oculta': '1'})['incluir_legend']
    assert 'rep_gantt_leyenda_oculta' in pr.FORMATO_CLAVES
    # El diálogo la lee y la escribe
    import inspect
    from views import formato_reporte_dialog as FRD
    src = inspect.getsource(FRD)
    assert "rep_gantt_leyenda_oculta" in src and "chk_gantt_leyenda" in src
    # Y `_render_pdf_completo` ya no fuerza True: sus cuatro flags aceptan None
    from views.cronograma_view import GanttWidget
    sig = inspect.signature(GanttWidget._render_pdf_completo)
    assert all(sig.parameters[k].default is None
               for k in ('incluir_header', 'incluir_footer', 'incluir_legend', 'incluir_page'))


def test_los_editables_obedecen_a_encabezado_y_pie_apagados():
    """Marco, 9 sep 2026: con «Imprimir el encabezado/pie» apagados, Excel,
    ODS, Word y ODT seguían llevándolos. Los cuatro salen de helpers comunes
    (`_xlsx_header_pdf_style`/`_xlsx_encabezado` + `_setup_impresion` en
    Excel —el ODS se convierte del .xlsx—; `_add_header_marca`/`_add_footer`
    en Word —el ODT del .docx—), así que la casilla se mira ahí. La tabla
    sube exactamente las filas del encabezado y no pierde ninguna."""
    import io
    import docx
    import openpyxl
    import core.exporter as EX
    import core.word_reports as WR

    def _flags(enc, pie):
        d.set_config('rep_encabezado_oculto', '1' if enc else '0')
        d.set_config('rep_pie_oculto', '1' if pie else '0')

    def _hoja():
        return openpyxl.load_workbook(io.BytesIO(EX.exportar_presupuesto(PID).getvalue())).worksheets[0]

    def _fila_tabla(ws):
        return next(r[0].row for r in ws.iter_rows(min_row=1, max_row=12)
                    if any(c.value and str(c.value).strip() == 'Ítem' for c in r))

    try:
        _flags(False, False)
        ws = _hoja()
        arriba = ' '.join(str(c.value) for row in ws.iter_rows(min_row=1, max_row=3) for c in row if c.value)
        assert 'Costo al' in arriba and ws.oddFooter.right.text
        f0, n0 = _fila_tabla(ws), ws.max_row
        _flags(True, True)
        ws = _hoja()
        arriba = ' '.join(str(c.value) for row in ws.iter_rows(min_row=1, max_row=3) for c in row if c.value)
        assert 'Costo al' not in arriba and 'IngePresupuestos' not in arriba
        assert not ws.oddFooter.right.text
        assert _fila_tabla(ws) == f0 - 3 and ws.max_row == n0 - 3     # sube 3 filas, ninguna se pierde
        # Word: cabecera y pie vacíos, mismo cuerpo
        def _word():
            fd, p = tempfile.mkstemp(suffix='_res.docx'); os.close(fd)
            _tmpfiles.append(p)
            WR.generar_word('resumen', PID, p)
            doc = docx.Document(p); sec = doc.sections[0]
            h = ''.join(c.text for t in sec.header.tables for r in t.rows for c in r.cells) + ''.join(p_.text for p_ in sec.header.paragraphs)
            f = ''.join(c.text for t in sec.footer.tables for r in t.rows for c in r.cells) + ''.join(p_.text for p_ in sec.footer.paragraphs)
            return h.strip(), f.strip(), len(doc.paragraphs), len(doc.tables)
        h1, f1, np1, nt1 = _word()
        assert not h1 and not f1
        _flags(False, False)
        h0, f0_, np0, nt0 = _word()
        assert h0 and f0_ and (np0, nt0) == (np1, nt1)
    finally:
        _flags(False, False)
